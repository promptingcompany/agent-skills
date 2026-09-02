#!/usr/bin/env python3
"""Print .xlsx / .csv as markdown or csv. No tpc dependency."""

from __future__ import annotations

import argparse
import csv
import io
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}
MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
CSV_ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")


def die(message: str, code: int = 1) -> None:
    print(message, file=sys.stderr)
    raise SystemExit(code)


def cell_col(ref: str) -> int:
    col = 0
    for ch in ref:
        if ch.isalpha():
            col = col * 26 + (ord(ch.upper()) - 64)
        else:
            break
    return max(col - 1, 0)


def escape_md(value: object) -> str:
    text = "" if value is None else str(value)
    return text.replace("|", "\\|").replace("\n", " ").replace("\r", "")


def render_md(rows: list[list[object]]) -> str:
    if not rows:
        return "_(empty)_"
    width = max(len(row) for row in rows)
    normalized = [list(row) + [""] * (width - len(row)) for row in rows]
    header, *body = normalized
    lines = [
        "| " + " | ".join(escape_md(c) for c in header) + " |",
        "|" + "|".join("---" for _ in header) + "|",
    ]
    for row in body:
        lines.append("| " + " | ".join(escape_md(c) for c in row) + " |")
    return "\n".join(lines)


def render_csv(rows: list[list[object]]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    for row in rows:
        writer.writerow(["" if c is None else str(c) for c in row])
    return buf.getvalue()


def trim_rows(rows: list[list[object]], max_rows: int) -> tuple[list[list[object]], int]:
    if max_rows <= 0 or len(rows) <= max_rows:
        return rows, 0
    header, body = rows[:1], rows[1:]
    kept = header + body[: max(max_rows - 1, 0)]
    return kept, len(rows) - len(kept)


# --- CSV -------------------------------------------------------------------

def _bom_encoding(data: bytes) -> str | None:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return "utf-16"
    if data.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    return None


def _csv_encodings(data: bytes) -> list[str]:
    bom = _bom_encoding(data)
    encodings: list[str] = []
    if bom:
        encodings.append(bom)
    # utf-16 without a BOM still has NUL bytes between ASCII characters
    if bom != "utf-16" and b"\x00" in data[:512]:
        encodings.append("utf-16")
    for encoding in CSV_ENCODINGS:
        if encoding not in encodings:
            encodings.append(encoding)
    return encodings


def read_csv_bytes(data: bytes) -> tuple[list[list[object]], str]:
    last_error = None
    for encoding in _csv_encodings(data):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError as exc:
            last_error = exc
            continue
        if "\x00" in text:
            continue
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        return [list(row) for row in reader], encoding
    die(f"Could not decode CSV ({last_error})")


def read_csv_path(path: Path) -> tuple[list[list[object]], str, str]:
    return (*read_csv_bytes(path.read_bytes()), "stdlib")


# --- xlsx: pandas / openpyxl / stdlib --------------------------------------

def read_xlsx_pandas(path: Path) -> tuple[dict[str, list[list[object]]], str]:
    import pandas as pd  # type: ignore

    book = pd.ExcelFile(path)
    sheets: dict[str, list[list[object]]] = {}
    for name in book.sheet_names:
        frame = book.parse(name, header=None, dtype=object)
        sheets[str(name)] = [
            [("" if v != v else v) for v in row]  # NaN -> ""
            for row in frame.itertuples(index=False, name=None)
        ]
    return sheets, "pandas"


def read_xlsx_openpyxl(path: Path) -> tuple[dict[str, list[list[object]]], str]:
    from openpyxl import load_workbook  # type: ignore

    book = load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, list[list[object]]] = {}
    for worksheet in book.worksheets:
        rows = [[("" if c is None else c) for c in row] for row in worksheet.iter_rows(values_only=True)]
        sheets[worksheet.title] = rows
    return sheets, "openpyxl"


def _shared_strings(zf: zipfile.ZipFile) -> list[str]:
    try:
        root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    except KeyError:
        return []
    out: list[str] = []
    for si in root.findall(f"{MAIN}si"):
        out.append("".join((t.text or "") for t in si.iter(f"{MAIN}t")))
    return out


def _sheet_rows(xml: bytes, shared: list[str]) -> list[list[object]]:
    root = ET.fromstring(xml)
    rows: list[list[object]] = []
    for row_el in root.iter(f"{MAIN}row"):
        cells: dict[int, object] = {}
        max_idx = -1
        for cell in row_el.findall(f"{MAIN}c"):
            ref = cell.get("r") or ""
            idx = cell_col(ref)
            max_idx = max(max_idx, idx)
            kind = cell.get("t")
            value_el = cell.find(f"{MAIN}v")
            is_el = cell.find(f"{MAIN}is")
            if kind == "inlineStr" and is_el is not None:
                cells[idx] = "".join((t.text or "") for t in is_el.iter(f"{MAIN}t"))
            elif value_el is None or value_el.text is None:
                cells[idx] = ""
            elif kind == "s":
                try:
                    cells[idx] = shared[int(value_el.text)]
                except (ValueError, IndexError):
                    cells[idx] = value_el.text
            elif kind == "b":
                cells[idx] = "TRUE" if value_el.text == "1" else "FALSE"
            else:
                cells[idx] = value_el.text
        if max_idx < 0:
            rows.append([])
        else:
            rows.append([cells.get(i, "") for i in range(max_idx + 1)])
    return rows


def read_xlsx_stdlib(path: Path) -> tuple[dict[str, list[list[object]]], str]:
    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        die(f"Not a valid .xlsx zip: {path}")

    with zf:
        names = set(zf.namelist())
        if "xl/workbook.xml" not in names:
            die(f"Not a spreadsheet workbook: {path}")
        shared = _shared_strings(zf)
        workbook = ET.fromstring(zf.read("xl/workbook.xml"))
        rels_root = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        rels = {
            rel.get("Id"): rel.get("Target")
            for rel in rels_root.findall(f"{{{NS['pr']}}}Relationship")
        }
        sheets: dict[str, list[list[object]]] = {}
        for sheet in workbook.findall("m:sheets/m:sheet", NS):
            title = sheet.get("name") or "Sheet"
            rel_id = sheet.get(f"{{{NS['r']}}}id")
            target = rels.get(rel_id or "")
            if not target:
                sheets[title] = []
                continue
            member = target if target.startswith("xl/") else f"xl/{target.lstrip('/')}"
            if member not in names:
                sheets[title] = []
                continue
            sheets[title] = _sheet_rows(zf.read(member), shared)
    return sheets, "stdlib"


def read_xlsx(path: Path) -> tuple[dict[str, list[list[object]]], str]:
    errors: list[str] = []
    for loader in (read_xlsx_pandas, read_xlsx_openpyxl, read_xlsx_stdlib):
        try:
            return loader(path)
        except ImportError:
            continue
        except SystemExit:
            raise
        except Exception as exc:
            errors.append(f"{loader.__name__.removeprefix('read_xlsx_')}: {exc}")
    detail = "; ".join(errors) if errors else "no backend available"
    die(f"Could not read xlsx ({detail})")


def looks_like_xlsx(path: Path) -> bool:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return True
    if suffix in {".csv", ".tsv"}:
        return False
    try:
        return zipfile.is_zipfile(path) and "xl/workbook.xml" in zipfile.ZipFile(path).namelist()
    except OSError:
        return False


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Print a local .xlsx or .csv as a markdown table. Does not use tpc."
    )
    parser.add_argument("path", type=Path, help="Path to .xlsx or .csv")
    parser.add_argument("--max-rows", type=int, default=200, help="Max rows per sheet, including header")
    parser.add_argument("--format", choices=("markdown", "csv"), default="markdown")
    parser.add_argument("--sheet", help="Only this sheet name (xlsx)")
    parser.add_argument("--out", type=Path, help="Write output to this file instead of stdout")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(sys.argv[1:] if argv is None else argv)
    path = args.path.expanduser()
    if not path.is_file():
        die(f"File not found: {path}")

    chunks: list[str] = []
    if looks_like_xlsx(path):
        sheets, backend = read_xlsx(path)
        if args.sheet:
            if args.sheet not in sheets:
                die(f"Sheet not found: {args.sheet}. Available: {', '.join(sheets) or '(none)'}")
            sheets = {args.sheet: sheets[args.sheet]}
        if not sheets:
            die("Workbook has no sheets")
        for name, rows in sheets.items():
            shown, omitted = trim_rows(rows, args.max_rows)
            body = render_csv(shown) if args.format == "csv" else render_md(shown)
            chunks.append(f"## {name}\n\n{body}\n")
            chunks.append(
                f"_backend={backend}; rows_shown={len(shown)}; rows_omitted={omitted}_\n"
            )
    else:
        rows, encoding, backend = read_csv_path(path)
        shown, omitted = trim_rows(rows, args.max_rows)
        body = render_csv(shown) if args.format == "csv" else render_md(shown)
        chunks.append(body + "\n")
        chunks.append(
            f"_backend={backend}; encoding={encoding}; rows_shown={len(shown)}; rows_omitted={omitted}_\n"
        )

    output = "\n".join(chunks).rstrip() + "\n"
    if args.out:
        args.out.write_text(output, encoding="utf-8")
        print(f"Wrote {args.out}", file=sys.stderr)
    else:
        sys.stdout.write(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
